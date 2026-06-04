"""
Importa inventario desde Control_Inventario_Marmoleria *.xlsx

Hojas:
  - "Control de Bloques" → placas + lotes (código IMPORT-…)
  - "Inventario Remanentes" → retazos (hasta la sección INSUMOS)

Columnas Bloques: Fecha | Bloque | Descripcion Material | Largo | Alto | m2 | Valor m2 | Observaciones
Remanentes: Fecha | Material/Color | Espesor (mm) | Tipo corte | Largo (mm) | Ancho (mm) | m2 | Estado | … | Bloque | Obs.

Uso (desde carpeta backend, venv activado):
  .venv\\Scripts\\python scripts/import_control_inventario_xlsx.py --sqlite --dry-run
  .venv\\Scripts\\python scripts/import_control_inventario_xlsx.py --sqlite --wipe-inventory --create-materials
  .venv\\Scripts\\python scripts/import_control_inventario_xlsx.py --wipe-inventory --wipe-lotes --create-materials

Con PostgreSQL (.env): omitir --sqlite.
"""

from __future__ import annotations

import argparse
import os
import re
import sys
import unicodedata
from datetime import datetime
from pathlib import Path
from typing import Any, Callable

_BACKEND = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(_BACKEND))

from sqlalchemy import delete, or_  # noqa: E402
from sqlalchemy.engine import Engine  # noqa: E402
from sqlalchemy.orm import Session, sessionmaker  # noqa: E402

from app.models import Lote, Material, MovimientoInventario, Placa, Retazo  # noqa: E402


def _default_xlsx_path() -> Path:
    """Repo: sube directorios hasta encontrar el .xlsx. En imagen Docker suele no existir: usar --file."""
    here = Path(__file__).resolve()
    name = "Control_Inventario_Marmoleria 2026.xlsx"
    for i in range(0, len(here.parents)):
        cand = here.parents[i] / name
        if cand.is_file():
            return cand
    raise SystemExit(
        "No se encontró el .xlsx por defecto; montá el archivo y usá --file /ruta/dentro/del/contenedor"
    )


def _default_sqlite_path() -> Path:
    return _BACKEND.parent / "dev_inventory.db"


def _norm_txt(s: str) -> str:
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode("ascii")
    return " ".join(s.lower().split())


def _parse_mm_from_material(text: str | None) -> int | None:
    if not text:
        return None
    m = re.search(r"(\d{1,2})\s*mm", text, re.I)
    if m:
        return int(m.group(1))
    return None


def _to_int_mm(val: Any) -> int | None:
    if val is None:
        return None
    if isinstance(val, (int, float)):
        if val <= 0:
            return None
        return int(round(float(val)))
    s = str(val).strip().replace(",", ".")
    if not s:
        return None
    try:
        v = float(s)
        return int(round(v)) if v > 0 else None
    except ValueError:
        return None


def _float_or_zero(val: Any) -> float:
    if isinstance(val, (int, float)):
        return float(val)
    if val is None or not str(val).strip():
        return 0.0
    try:
        return float(str(val).replace(",", "."))
    except ValueError:
        return 0.0


def _load_rows_control_bloques(path: Path, sheet: str) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        wb.close()
        raise SystemExit(f"Hoja no encontrada: {sheet!r}. Disponibles: {wb.sheetnames}")
    ws = wb[sheet]
    rows_iter = ws.iter_rows(values_only=True)
    header_idx: int | None = None
    header_map: dict[str, int] = {}
    for i, row in enumerate(rows_iter):
        if not row:
            continue
        cells = [str(c).strip() if c is not None else "" for c in row]
        joined = " ".join(cells).lower()
        if "descripcion material" in joined and "bloque" in joined and "largo" in joined:
            header_idx = i
            for j, c in enumerate(row):
                if c is None:
                    continue
                key = str(c).strip().lower().replace("\n", " ")
                if "fecha" in key and "insumo" not in key:
                    header_map["fecha"] = j
                elif key == "bloque" or key.startswith("bloque"):
                    header_map["bloque"] = j
                elif "descripcion" in key or "material" in key:
                    header_map["material"] = j
                elif "largo" in key:
                    header_map["largo"] = j
                elif "alto" in key:
                    header_map["alto"] = j
                elif "m2" in key and "valor" not in key:
                    header_map["m2"] = j
                elif "valor" in key and "m2" in key:
                    header_map["valor_m2"] = j
                elif "observ" in key:
                    header_map["obs"] = j
            break
    if header_idx is None:
        wb.close()
        raise SystemExit("No se encontró la fila de encabezados en Control de Bloques.")

    out: list[dict[str, Any]] = []

    def get(r: tuple[Any, ...], name: str) -> Any:
        idx = header_map.get(name)
        if idx is None or idx >= len(r):
            return None
        return r[idx]

    for row in ws.iter_rows(min_row=header_idx + 2, values_only=True):
        if not row:
            continue
        material_txt = get(row, "material")
        if material_txt is None or str(material_txt).strip() == "":
            continue
        li = _to_int_mm(get(row, "largo"))
        wi = _to_int_mm(get(row, "alto"))
        if li is None or wi is None:
            continue
        bloque = get(row, "bloque")
        out.append(
            {
                "fecha": get(row, "fecha"),
                "bloque": str(bloque).strip() if bloque is not None else "",
                "material": str(material_txt).strip(),
                "largo_mm": li,
                "alto_mm": wi,
                "m2": get(row, "m2"),
                "valor_m2": get(row, "valor_m2"),
                "obs": get(row, "obs"),
            }
        )
    wb.close()
    return out


def _load_rows_remanentes(path: Path, sheet: str) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        wb.close()
        raise SystemExit(f"Hoja no encontrada: {sheet!r}. Disponibles: {wb.sheetnames}")
    ws = wb[sheet]
    header_idx: int | None = None
    header_map: dict[str, int] = {}
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if not row:
            continue
        cells = [str(c).strip() if c is not None else "" for c in row]
        joined = " ".join(cells).lower()
        if "material" in joined and "largo" in joined and "ancho" in joined:
            header_idx = i
            for j, c in enumerate(row):
                if c is None:
                    continue
                key = str(c).strip().lower().replace("\n", " ")
                if "fecha" in key and "insumo" not in key:
                    header_map["fecha"] = j
                elif "material" in key or "color" in key:
                    header_map["material"] = j
                elif "espesor" in key:
                    header_map["espesor"] = j
                elif "corte" in key or "martillo" in key:
                    header_map["corte"] = j
                elif "largo" in key:
                    header_map["largo"] = j
                elif "ancho" in key:
                    header_map["ancho"] = j
                elif key == "m2" or (("m2" in key) and "valor" not in key):
                    header_map["m2"] = j
                elif "estado" in key:
                    header_map["estado"] = j
                elif "bloque" in key:
                    header_map["bloque"] = j
                elif "observ" in key:
                    header_map["obs"] = j
            break
    if header_idx is None:
        wb.close()
        raise SystemExit("No se encontró encabezado en Inventario Remanentes.")

    out: list[dict[str, Any]] = []

    def get(r: tuple[Any, ...], name: str) -> Any:
        idx = header_map.get(name)
        if idx is None or idx >= len(r):
            return None
        return r[idx]

    for row in ws.iter_rows(min_row=header_idx + 2, values_only=True):
        if not row:
            continue
        c1 = str(row[1]).strip().upper() if len(row) > 1 and row[1] is not None else ""
        c2 = str(row[2]).strip().upper() if len(row) > 2 and row[2] is not None else ""
        if c1 == "FECHA" and "INSUMO" in c2:
            break

        material_txt = get(row, "material")
        if material_txt is None or str(material_txt).strip() == "":
            continue
        li = _to_int_mm(get(row, "largo"))
        wi = _to_int_mm(get(row, "ancho"))
        if li is None or wi is None:
            continue

        esp_raw = get(row, "espesor")
        esp_i = _to_int_mm(esp_raw) if esp_raw is not None else None

        bloque = get(row, "bloque")
        out.append(
            {
                "fecha": get(row, "fecha"),
                "material": str(material_txt).strip(),
                "espesor_mm": esp_i,
                "corte": get(row, "corte"),
                "largo_mm": li,
                "ancho_mm": wi,
                "m2_excel": get(row, "m2"),
                "estado_excel": get(row, "estado"),
                "bloque": str(bloque).strip() if bloque is not None else "",
                "obs": get(row, "obs"),
            }
        )
    wb.close()
    return out


def _find_material(db: Session, nombre: str) -> Material | None:
    n = nombre.strip()
    if not n:
        return None
    m = db.query(Material).filter(Material.nombre == n).first()
    if m:
        return m
    cand = (
        db.query(Material)
        .filter(or_(Material.nombre.ilike(f"%{n[:50]}%"), Material.nombre.ilike(f"%{n[:25]}%")))
        .first()
    )
    if cand:
        return cand
    parts = [p for p in re.split(r"\W+", n) if len(p) > 3]
    for p in parts[:3]:
        cand = db.query(Material).filter(Material.nombre.ilike(f"%{p}%")).first()
        if cand:
            return cand
    return None


def _get_or_create_lote(
    db: Session,
    *,
    material_id: str,
    codigo_lote: str,
    ubicacion: str | None,
    notas: str | None,
) -> Lote:
    existing = db.query(Lote).filter(Lote.codigo_lote == codigo_lote).first()
    if existing:
        return existing
    lote = Lote(
        material_id=material_id,
        codigo_lote=codigo_lote,
        fecha_ingreso=datetime.now().strftime("%Y-%m-%d"),
        cantidad_inicial=0.0,
        stock_actual=0.0,
        cantidad=1,
        ubicacion=ubicacion,
        notas=notas,
        costo_m2=0.0,
        precio_venta=0.0,
        precio_mayorista=0.0,
    )
    db.add(lote)
    db.flush()
    return lote


def _wipe_inventory(session: Session, *, wipe_lotes: bool) -> None:
    session.execute(delete(MovimientoInventario))
    session.execute(delete(Placa))
    session.execute(delete(Retazo))
    if wipe_lotes:
        session.execute(delete(Lote))
    session.flush()


def _make_session_factory(args: argparse.Namespace) -> Callable[[], Session]:
    if args.sqlite:
        db_path = Path(args.sqlite_path).resolve()
        url = f"sqlite:///{db_path.as_posix()}"
        from sqlalchemy import create_engine

        eng: Engine = create_engine(url, echo=False)
        return sessionmaker(bind=eng, autoflush=False, autocommit=False)
    from app.db import SessionLocal

    return SessionLocal


def _import_bloques(
    session: Session,
    rows: list[dict[str, Any]],
    *,
    create_materials: bool,
    espesor_default: int,
) -> tuple[int, int]:
    created_p = created_m = 0
    for i, r in enumerate(rows, 1):
        bloque_raw = r["bloque"] or "SIN-BLOQUE"
        bloque_key = re.sub(r"\s+", "", str(bloque_raw)).upper()
        nombre_mat = r["material"]
        esp = _parse_mm_from_material(nombre_mat) or espesor_default

        mat = _find_material(session, nombre_mat)
        if mat is None and create_materials:
            vm2 = r["valor_m2"]
            precio = _float_or_zero(vm2)
            mat = Material(
                nombre=nombre_mat[:255],
                precio_m2=precio,
                espesor_mm=esp,
                stock_actual=0.0,
                unidad="m²",
                stock_minimo=0.0,
                ultima_actualizacion=datetime.now().strftime("%Y-%m-%d"),
                activo=True,
                disponible_para_venta=True,
            )
            session.add(mat)
            session.flush()
            created_m += 1
        elif mat is None:
            print(f"[omitido bloque] Sin material en DB: {nombre_mat[:70]}")
            continue

        assert mat is not None
        codigo_lote = f"IMPORT-{bloque_key}-{_norm_txt(mat.nombre)[:24].upper().replace(' ', '-')}"[:64]
        lote = _get_or_create_lote(
            session,
            material_id=mat.id,
            codigo_lote=codigo_lote,
            ubicacion=bloque_key,
            notas="Import Excel Control de Bloques",
        )

        codigo_placa = f"PLC-{bloque_key}-{i:04d}"[:64]
        if session.query(Placa).filter(Placa.codigo == codigo_placa).first():
            codigo_placa = f"PLC-{bloque_key}-{i:04d}-{os.urandom(2).hex()}"[:64]

        vm2 = r["valor_m2"]
        precio_placa = _float_or_zero(vm2) if _float_or_zero(vm2) > 0 else None

        placa = Placa(
            material_id=mat.id,
            lote_id=lote.id,
            largo=r["largo_mm"],
            ancho=r["alto_mm"],
            espesor=esp,
            codigo=codigo_placa,
            ubicacion=bloque_key,
            estado="disponible",
            precio=precio_placa,
        )
        session.add(placa)
        created_p += 1
    return created_m, created_p


def _import_remanentes(
    session: Session,
    rows: list[dict[str, Any]],
    *,
    create_materials: bool,
    espesor_default: int,
) -> tuple[int, int]:
    created_r = created_m = 0
    for i, r in enumerate(rows, 1):
        nombre_mat = r["material"]
        esp = r["espesor_mm"] or _parse_mm_from_material(nombre_mat) or espesor_default
        mat = _find_material(session, nombre_mat)
        if mat is None and create_materials:
            mat = Material(
                nombre=nombre_mat[:255],
                precio_m2=0.0,
                espesor_mm=int(esp),
                stock_actual=0.0,
                unidad="m²",
                stock_minimo=0.0,
                ultima_actualizacion=datetime.now().strftime("%Y-%m-%d"),
                activo=True,
                disponible_para_venta=True,
            )
            session.add(mat)
            session.flush()
            created_m += 1
        elif mat is None:
            print(f"[omitido remanente] Sin material en DB: {nombre_mat[:70]}")
            continue

        assert mat is not None
        bloque_raw = r["bloque"] or ""
        bloque_key = re.sub(r"\s+", "", str(bloque_raw)).upper() if bloque_raw else "REM"
        ubic = bloque_key if bloque_key else "REMANENTES"

        codigo_lote = f"IMPORT-RTZ-{bloque_key}-{_norm_txt(mat.nombre)[:16].upper().replace(' ', '-')}"[:64]
        lote = _get_or_create_lote(
            session,
            material_id=mat.id,
            codigo_lote=codigo_lote,
            ubicacion=ubic,
            notas="Import Excel Inventario Remanentes",
        )

        ret = Retazo(
            material_id=mat.id,
            lote_id=lote.id,
            largo=r["largo_mm"],
            ancho=r["ancho_mm"],
            espesor=int(esp),
            ubicacion=ubic[:64] if ubic else None,
            estado="disponible",
            en_venta=False,
            precio=None,
        )

        session.add(ret)
        created_r += 1
    return created_m, created_r


def main() -> None:
    parser = argparse.ArgumentParser(description="Importar inventario desde Excel de control")
    parser.add_argument("--file", type=Path, default=None, help="Ruta al .xlsx (en Docker suele ser obligatorio)")
    parser.add_argument("--sheet-bloques", default="Control de Bloques", help="Hoja placas/bloques")
    parser.add_argument("--sheet-remanentes", default="Inventario Remanentes", help="Hoja retazos")
    parser.add_argument("--dry-run", action="store_true", help="Solo leer Excel, no escribir BD")
    parser.add_argument("--sqlite", action="store_true", help="Usar SQLite (ver --sqlite-path)")
    parser.add_argument("--sqlite-path", type=Path, default=_default_sqlite_path(), help="Archivo .db para --sqlite")
    parser.add_argument("--wipe-inventory", action="store_true", help="Borrar movimientos, placas y retazos antes de importar")
    parser.add_argument("--wipe-lotes", action="store_true", help="También borrar todos los lotes (tras wipe-inventory)")
    parser.add_argument("--create-materials", action="store_true", help="Crear materiales faltantes desde el Excel")
    parser.add_argument("--espesor-default", type=int, default=20, help="Espesor mm por defecto")
    parser.add_argument("--no-bloques", action="store_true", help="No importar hoja Control de Bloques")
    parser.add_argument("--no-remanentes", action="store_true", help="No importar hoja Inventario Remanentes")
    args = parser.parse_args()
    if args.file is None:
        args.file = _default_xlsx_path()

    if not args.file.is_file():
        raise SystemExit(f"No existe el archivo: {args.file}")

    rows_b: list[dict[str, Any]] = []
    rows_r: list[dict[str, Any]] = []
    if not args.no_bloques:
        rows_b = _load_rows_control_bloques(args.file, args.sheet_bloques)
    if not args.no_remanentes:
        rows_r = _load_rows_remanentes(args.file, args.sheet_remanentes)

    print(f"Filas válidas — Bloques: {len(rows_b)} | Remanentes: {len(rows_r)}")

    if args.dry_run:
        for i, r in enumerate(rows_b[:15], 1):
            esp = _parse_mm_from_material(r["material"]) or args.espesor_default
            print(f"  B {i:3d} | {r['bloque']!s:6} | {r['largo_mm']}x{r['alto_mm']} mm | esp~{esp} | {r['material'][:55]!s}")
        if len(rows_b) > 15:
            print(f"  … (+{len(rows_b) - 15} bloques)")
        for i, r in enumerate(rows_r[:15], 1):
            esp = r["espesor_mm"] or _parse_mm_from_material(r["material"]) or args.espesor_default
            print(f"  R {i:3d} | {r['bloque']!s:6} | {r['largo_mm']}x{r['ancho_mm']} mm | esp~{esp} | {r['material'][:55]!s}")
        if len(rows_r) > 15:
            print(f"  … (+{len(rows_r) - 15} remanentes)")
        print("\nDry-run: no se escribió nada en la base.")
        return

    SessionFactory = _make_session_factory(args)
    if args.sqlite:
        print(f"[INFO] SQLite: {args.sqlite_path.resolve()}")

    session = SessionFactory()
    new_m_b = new_p = new_m_r = new_rz = 0
    try:
        if args.wipe_inventory:
            _wipe_inventory(session, wipe_lotes=args.wipe_lotes)
            print("Pendiente borrar en esta transacción: placas/retazos/movimientos" + (" + lotes" if args.wipe_lotes else "") + ".")

        if rows_b:
            new_m_b, new_p = _import_bloques(
                session, rows_b, create_materials=args.create_materials, espesor_default=args.espesor_default
            )
        if rows_r:
            new_m_r, new_rz = _import_remanentes(
                session, rows_r, create_materials=args.create_materials, espesor_default=args.espesor_default
            )

        session.commit()
        print(
            f"Listo. Materiales nuevos (bloques/rem): {new_m_b + new_m_r}. "
            f"Placas: {new_p}. Retazos: {new_rz}."
        )
    except Exception:
        session.rollback()
        raise
    finally:
        session.close()


if __name__ == "__main__":
    main()

